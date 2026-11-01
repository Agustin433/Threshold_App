"""Reglas de validez de evaluaciones con los casos reales de Supabase (11 tomas).

Cada regla vive en una sola funcion (modules/evaluation_rules.py) y se
verifica en tres capas:
- la funcion pura;
- el frame preparado (`_prepare_jump_df`), que consumen todas las vistas;
- la salida final: dashboard real (AppTest sobre app.py) y texto del PDF.

Casos reales (fila cruda de `evaluations`):
- Fernando Habib 30/10: pretension 1284.67 N, BW 90.24 kg (+399 N), RFD50 196 N/s.
- Alejandro Anadon 07/05: pretension 1167 N, BW 74 kg (+441 N), RFD50 389 N/s.
- Santiago Gray 07/05: bw_kg = 0.
- Ignacio De La Vega: CMJ 41.1 (15/05) -> 46.04 (31/10); el umbral "0.6 (Fijo)"
  salia de 1.5% de 41.1.
"""

from __future__ import annotations

import os
import re
import shutil
import unicodedata
import unittest
import uuid
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path

import numpy as np
import pandas as pd

from modules import evaluation_rules as rules
from modules.app_clock import frozen_app_now
from modules.force_time_analysis import interpret_imtp_force_time, summarize_force_time_test
from modules.jump_analysis import (
    _prepare_jump_df,
    build_composite_profile_snapshot,
    build_jump_delta_display_table,
    build_neuromuscular_profile_result,
    compute_swc_delta,
)
from modules.remote_store import jump_df_to_db_records

NOW = pd.Timestamp("2026-11-01 10:00")
APP_PY_PATH = Path(__file__).resolve().parent.parent / "app.py"
TEST_TMP_ROOT = Path(__file__).resolve().parent / ".tmp_evaluation_rules"

HABIB = {
    "Athlete": "Fernando Habib", "Date": "2026-10-30", "Source": "platform", "BW_kg": 90.24,
    "CMJ_cm": 30.44, "IMTP_N": 3883.43, "IMTP_pretension": 1284.67,
    "IMTP_rfd_50_N_s": 196.2, "IMTP_rfd_100_N_s": 0.0, "IMTP_rfd_150_N_s": 0.0, "IMTP_rfd_250_N_s": 200.0,
    "IMTP_force_50_N": 1294.48, "IMTP_force_100_N": 1284.67, "IMTP_force_150_N": 1284.67,
    "IMTP_force_200_N": 1284.67, "IMTP_force_250_N": 1333.7,
    "IMTP_force_L_N": 1900.0, "IMTP_force_R_N": 1983.0, "IMTP_asym_pct": 4.2,
    "CMJ_propulsive_PF_N": 2100.0,
}
ANADON = {
    "Athlete": "Alejandro Anadon", "Date": "2026-05-07", "Source": "platform", "BW_kg": 74.0,
    "CMJ_cm": 32.43, "SJ_cm": 31.63, "IMTP_N": 2526.0, "IMTP_pretension": 1167.0,
    "IMTP_rfd_50_N_s": 389.0, "IMTP_rfd_100_N_s": 1026.0,
    "IMTP_force_50_N": 1187.0, "IMTP_force_100_N": 1270.0, "IMTP_force_200_N": 1748.0, "IMTP_force_250_N": 1868.0,
}
GRAY = {
    "Athlete": "Santiago Gray", "Date": "2026-05-07", "Source": "platform", "BW_kg": 0.0,
    "CMJ_cm": 43.43, "SJ_cm": 44.0, "IMTP_N": 2705.0, "IMTP_pretension": 942.0,
    "IMTP_rfd_50_N_s": 8074.0, "IMTP_rfd_100_N_s": 6414.0,
    "IMTP_force_50_N": 1346.0, "IMTP_force_100_N": 1583.0, "IMTP_force_200_N": 1864.0, "IMTP_force_250_N": 1852.0,
}
DE_LA_VEGA = [
    {"Athlete": "Ignacio De La Vega", "Date": "2026-05-15", "Source": "platform", "BW_kg": 87.0, "CMJ_cm": 41.1},
    {
        "Athlete": "Ignacio De La Vega", "Date": "2026-10-31", "Source": "platform", "BW_kg": 87.5,
        "CMJ_cm": 46.04, "SJ_cm": 43.55, "IMTP_N": 3491.17, "IMTP_pretension": 1095.073,
        "IMTP_rfd_50_N_s": 12356.4,
    },
]
# Toma limpia (control): mismos datos que Habib con pretension = BW*g + 50 N.
CLEAN = dict(
    HABIB,
    Athlete="Control Limpio",
    IMTP_pretension=round(90.24 * 9.81 + 50, 2),
    IMTP_rfd_50_N_s=5300.0,
    IMTP_force_100_N=1800.0,
)

FLAG_TEXT = "pretensión alta: RFD y fuerza a 50–200 ms no interpretables"


def real_jump_df() -> pd.DataFrame:
    return pd.DataFrame([HABIB, ANADON, GRAY, *DE_LA_VEGA, CLEAN])


def _row(prepared: pd.DataFrame, athlete: str, date: str | None = None) -> pd.Series:
    data = prepared[prepared["Athlete"] == athlete]
    if date:
        data = data[data["Date"] == pd.Timestamp(date)]
    return data.iloc[-1]


def _ascii(text: str) -> str:
    return " ".join(unicodedata.normalize("NFD", text).encode("ascii", "ignore").decode("ascii").lower().split())


# ── 1. Calidad IMTP ────────────────────────────────────────────────────────
class ImtpQualityRuleTest(unittest.TestCase):
    def test_real_cases_are_flagged_with_both_reasons(self):
        habib = rules.assess_imtp_quality(1284.67, 90.24, 196.2)
        self.assertTrue(habib["flagged"])
        self.assertEqual(habib["label"], FLAG_TEXT)
        self.assertAlmostEqual(habib["pretension_excess_n"], 399.4, places=1)
        self.assertEqual(len(habib["reasons"]), 2)

        anadon = rules.assess_imtp_quality(1167, 74, 389)
        self.assertTrue(anadon["flagged"])
        self.assertAlmostEqual(anadon["pretension_excess_n"], 441.1, places=1)

    def test_each_criterion_flags_on_its_own(self):
        self.assertTrue(rules.assess_imtp_quality(90.24 * 9.81 + 201, 90.24, 5000)["flagged"])
        self.assertTrue(rules.assess_imtp_quality(90.24 * 9.81, 90.24, 999)["flagged"])
        self.assertFalse(rules.assess_imtp_quality(90.24 * 9.81 + 200, 90.24, 1000)["flagged"])

    def test_prepared_frame_drops_time_dependent_values_but_keeps_peak_and_relative(self):
        prepared = _prepare_jump_df(real_jump_df())
        for athlete in ("Fernando Habib", "Alejandro Anadon"):
            row = _row(prepared, athlete)
            self.assertEqual(row[rules.IMTP_QUALITY_FLAG_COLUMN], rules.IMTP_HIGH_PRETENSION_FLAG)
            for column in rules.IMTP_TIME_DEPENDENT_COLUMNS:
                if column in prepared.columns:
                    self.assertTrue(pd.isna(row[column]), f"{athlete} {column}")
            self.assertTrue(pd.notna(row["IMTP_N"]))
            self.assertTrue(pd.notna(row["IMTP_relPF"]))
        clean = _row(prepared, "Control Limpio")
        self.assertTrue(pd.isna(clean[rules.IMTP_QUALITY_FLAG_COLUMN]))
        self.assertEqual(float(clean["IMTP_rfd_50_N_s"]), 5300.0)

    def test_flag_survives_a_second_prepare(self):
        twice = _prepare_jump_df(_prepare_jump_df(real_jump_df()))
        self.assertEqual(_row(twice, "Fernando Habib")[rules.IMTP_QUALITY_FLAG_COLUMN], rules.IMTP_HIGH_PRETENSION_FLAG)

    def test_dsi_uses_peak_force_only(self):
        row = _row(_prepare_jump_df(real_jump_df()), "Fernando Habib")
        self.assertAlmostEqual(float(row["DSI"]), round(2100.0 / 3883.43, 3), places=3)

    def test_force_time_summary_hides_points_and_keeps_row_selectable(self):
        summary = summarize_force_time_test(HABIB, test_id="imtp")
        self.assertEqual(summary["quality_label"], FLAG_TEXT)
        self.assertTrue(summary["time_points_suppressed"])
        self.assertIsNone(summary["rfd_50_n_s"])
        self.assertIsNone(summary["force_100_n"])
        self.assertEqual(summary["peak_force_n"], 3883.43)
        interpretation = interpret_imtp_force_time(summary)
        self.assertEqual(interpretation["quality_label"], FLAG_TEXT)
        self.assertIn("no interpretable", interpretation["rfd_text"])

    def test_raw_values_are_not_erased_on_save(self):
        prepared = _prepare_jump_df(pd.DataFrame([HABIB]))
        record = jump_df_to_db_records(prepared)[0]
        self.assertNotIn("imtp_rfd_50_n_s", record)  # NaN no se escribe: el crudo queda en la base


# ── 2. Fuerza relativa bruta y neta ────────────────────────────────────────
class ImtpRelativeForceTest(unittest.TestCase):
    def test_gross_and_net_values(self):
        gross, net = rules.imtp_relative_force(3883.43, 90.24)
        self.assertAlmostEqual(gross, 43.03, places=2)
        self.assertAlmostEqual(net, round((3883.43 - 90.24 * 9.81) / 90.24, 2), places=2)
        self.assertAlmostEqual(gross - net, 9.81, places=1)

    def test_prepared_frame_has_both_columns(self):
        row = _row(_prepare_jump_df(real_jump_df()), "Fernando Habib")
        self.assertAlmostEqual(float(row["IMTP_relPF"]), 43.03, places=2)
        self.assertAlmostEqual(float(row["IMTP_relPF_net"]), 33.22, places=2)


# ── 3. Peso corporal invalido ─────────────────────────────────────────────
class BodyWeightRuleTest(unittest.TestCase):
    def test_validation_range(self):
        self.assertEqual(rules.validate_bw_kg_input(None), (None, None))
        self.assertEqual(rules.validate_bw_kg_input(74), (74.0, None))
        for bad in (0, 29.9, 200.1, -5):
            value, warning = rules.validate_bw_kg_input(bad)
            self.assertIsNone(value)
            self.assertIn("fuera de rango", warning)

    def test_santiago_gray_zero_weight_blocks_relative_metrics(self):
        row = _row(_prepare_jump_df(real_jump_df()), "Santiago Gray")
        self.assertTrue(pd.isna(row["BW_kg"]))
        self.assertTrue(pd.isna(row["IMTP_relPF"]))
        self.assertTrue(pd.isna(row["IMTP_relPF_net"]))
        self.assertTrue(pd.isna(row.get("Jump_Momentum")))
        self.assertEqual(float(row["IMTP_N"]), 2705.0)

    def test_previously_computed_relative_metrics_are_cleared(self):
        frame = pd.DataFrame([dict(GRAY, IMTP_relPF=999.0, Jump_Momentum=1.0)])
        row = _prepare_jump_df(frame).iloc[0]
        self.assertTrue(pd.isna(row["IMTP_relPF"]))
        self.assertTrue(pd.isna(row["Jump_Momentum"]))

    def test_zero_weight_is_saved_as_explicit_null(self):
        record = jump_df_to_db_records(pd.DataFrame([GRAY]))[0]
        self.assertIn("bw_kg", record)
        self.assertIsNone(record["bw_kg"])

    def test_migration_nulls_and_constrains_bw(self):
        sql = (Path(__file__).resolve().parent.parent / "supabase/migrations/20261101120000_evaluations_bw_kg_range.sql").read_text()
        self.assertIn("set bw_kg = null", sql)
        self.assertIn("bw_kg >= 30 and bw_kg <= 200", sql)


# ── 4. Antiguedad de la evaluacion ────────────────────────────────────────
class StaleEvaluationRuleTest(unittest.TestCase):
    def test_weeks_and_label(self):
        result = rules.evaluation_staleness("2026-05-07", NOW)
        self.assertTrue(result["is_stale"])
        self.assertEqual(result["weeks"], 25)
        self.assertEqual(result["label"], "Evaluación desactualizada (hace 25 semanas)")
        self.assertFalse(rules.evaluation_staleness("2026-10-30", NOW)["is_stale"])
        self.assertFalse(rules.evaluation_staleness(NOW - pd.Timedelta(weeks=12), NOW)["is_stale"])

    def test_composite_profile_flags_anadon_and_lowers_confidence(self):
        prepared = _prepare_jump_df(real_jump_df())
        snapshot, _ = build_composite_profile_snapshot(prepared[prepared["Athlete"] == "Alejandro Anadon"])
        stale = build_neuromuscular_profile_result(snapshot, context={"reference_date": NOW})
        fresh = build_neuromuscular_profile_result(snapshot, context={"reference_date": "2026-05-10"})
        self.assertIn("stale_evaluation", stale["flags"])
        self.assertEqual(stale["staleness_label"], "Evaluación desactualizada (hace 25 semanas)")
        self.assertIn("Evaluación desactualizada (hace 25 semanas)", stale["profile_source_note"])
        self.assertEqual(stale["confidence"], rules.downgrade_confidence(fresh["confidence"]))
        if fresh["confidence"] != "low":
            self.assertNotEqual(stale["confidence"], fresh["confidence"])

    def test_downgrade(self):
        self.assertEqual(rules.downgrade_confidence("high"), "moderate")
        self.assertEqual(rules.downgrade_confidence("moderate"), "low")
        self.assertEqual(rules.downgrade_confidence("low"), "low")


# ── 5. Umbral de cambio ───────────────────────────────────────────────────
class ChangeThresholdRuleTest(unittest.TestCase):
    def test_de_la_vega_cmj_uses_instrument_typical_error(self):
        delta = compute_swc_delta(_prepare_jump_df(pd.DataFrame(DE_LA_VEGA)), "2026-10-31")
        cmj = delta[delta["Variable"] == "CMJ_cm"].iloc[0]
        # Antes: 41.1 * 1.5% = 0.62 "(Fijo)", por debajo del error de medicion.
        self.assertAlmostEqual(float(cmj["Threshold_abs"]), 41.1 * 0.03, places=3)
        self.assertEqual(cmj["Threshold_method"], "TE instrumento")
        self.assertEqual(cmj["Threshold_label"], "TE instrumento 3%")
        display = build_jump_delta_display_table(delta)
        text = display.loc[display["Variable"] == "CMJ", "Threshold"].iloc[0]
        self.assertEqual(text, "1.2 (TE instrumento 3%)")
        self.assertNotIn("Fijo", " ".join(display["Threshold"].astype(str)))

    def test_swc_wins_when_larger_than_te(self):
        result = rules.resolve_change_threshold("CMJ_cm", 40.0, [40.0, 42.0], swc_pct=5.0)
        self.assertEqual(result["method"], "SWC fijo")
        self.assertAlmostEqual(result["threshold_abs"], 2.0)

    def test_intraindividual_from_three_measurements(self):
        result = rules.resolve_change_threshold("CMJ_cm", 40.0, [40.0, 44.0, 39.0], swc_pct=1.5)
        sd = float(np.std([40.0, 44.0, 39.0], ddof=1))
        te = float(np.std(np.diff([40.0, 44.0, 39.0]), ddof=1) / np.sqrt(2))
        self.assertAlmostEqual(result["threshold_abs"], max(0.2 * sd, te))
        self.assertTrue(result["label"].endswith("N=3"))


# ── Mismo resultado en dashboard (AppTest) y PDF ──────────────────────────
@contextmanager
def isolated_store():
    original_store = os.environ.get("THRESHOLD_STORE_DIR")
    original_supabase = {key: os.environ.pop(key, None) for key in ("SUPABASE_URL", "SUPABASE_KEY", "SUPABASE_SERVICE_ROLE_KEY")}
    tmp_root = TEST_TMP_ROOT / f"rules_{uuid.uuid4().hex[:8]}"
    try:
        store_dir = tmp_root / "store"
        store_dir.mkdir(parents=True, exist_ok=True)
        os.environ["THRESHOLD_STORE_DIR"] = str(store_dir)
        import importlib

        import local_store

        importlib.reload(local_store)
        yield store_dir
    finally:
        if original_store is None:
            os.environ.pop("THRESHOLD_STORE_DIR", None)
        else:
            os.environ["THRESHOLD_STORE_DIR"] = original_store
        for key, value in original_supabase.items():
            if value is not None:
                os.environ[key] = value
        shutil.rmtree(tmp_root, ignore_errors=True)
        import importlib

        import local_store

        importlib.reload(local_store)
        if TEST_TMP_ROOT.exists() and not any(TEST_TMP_ROOT.iterdir()):
            TEST_TMP_ROOT.rmdir()


def _dashboard_text(at) -> str:
    parts: list[str] = []
    for element_type in ("markdown", "caption", "error", "warning", "info", "metric"):
        for element in getattr(at, element_type, []):
            parts.append(str(getattr(element, "value", "")))
            parts.append(str(getattr(element, "label", "")))
    for frame in getattr(at, "dataframe", []):
        try:
            parts.append(frame.value.to_string())
        except Exception:
            pass
    return " ".join(parts)


def _render_platform_dashboard(athlete: str) -> str:
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_file(str(APP_PY_PATH), default_timeout=180)
    at.run()
    assert not at.exception, at.exception
    nav = next(bg for bg in at.button_group if bg.label == "Vista principal")
    nav.set_value("Análisis Plataforma").run()
    assert not at.exception, at.exception
    if not any(sb.key == "ev_jugador" for sb in at.selectbox):
        raise AssertionError("sin selector de atleta: " + _dashboard_text(at)[:800])
    at.selectbox(key="ev_jugador").set_value(athlete).run()
    assert not at.exception, at.exception
    return _dashboard_text(at)


def _pdf_text(pdf: bytes) -> str:
    from pypdf import PdfReader

    from modules import report_generator

    pages = [(page.extract_text() or "") for page in PdfReader(BytesIO(pdf)).pages]
    return _ascii(" ".join(report_generator._repair_mojibake_text(page) for page in pages))


def _pdf_state() -> dict:
    return {"jump_df": _prepare_jump_df(real_jump_df()), "athletes_list": sorted(real_jump_df()["Athlete"].unique())}


class ExcelExportRulesTest(unittest.TestCase):
    def test_export_sheet_applies_all_imtp_and_weight_rules(self):
        from openpyxl import load_workbook

        from modules import report_generator

        with frozen_app_now(NOW):
            sheets = report_generator.build_report_sheets(_pdf_state(), "Todos", "profe", include_technical_annex=True)
            workbook = load_workbook(BytesIO(report_generator.export_excel(sheets)))
        sheet = workbook["07_Evaluaciones"]
        rows = list(sheet.iter_rows(values_only=True))
        header = list(rows[0])
        data = {row[header.index("Athlete")]: dict(zip(header, row)) for row in rows[1:]}
        habib = data["Fernando Habib"]
        self.assertEqual(habib["IMTP calidad de toma"], FLAG_TEXT)
        self.assertIsNone(habib.get("IMTP_rfd_50_N_s"))
        self.assertAlmostEqual(habib["IMTP relPF bruta (N/kg) = PF/masa"], 43.03, places=2)
        self.assertAlmostEqual(habib["IMTP relPF neta (N/kg) = (PF-BW)/masa"], 33.22, places=2)
        gray = data["Santiago Gray"]
        self.assertIsNone(gray["BW_kg"])
        self.assertIsNone(gray["IMTP relPF bruta (N/kg) = PF/masa"])


class DashboardAndPdfParityTest(unittest.TestCase):
    """La misma regla produce el mismo resultado visible en ambas superficies."""

    @classmethod
    def setUpClass(cls):
        from modules import report_generator

        cls.dashboard: dict[str, str] = {}
        cls.pdf: dict[tuple[str, str], str] = {}
        with frozen_app_now(NOW), isolated_store() as store_dir:
            real_jump_df().to_csv(store_dir / "evaluations_history.csv", index=False)
            for athlete in ("Fernando Habib", "Alejandro Anadon", "Santiago Gray", "Ignacio De La Vega"):
                cls.dashboard[athlete] = _ascii(_render_platform_dashboard(athlete))
            state = _pdf_state()
            for athlete in ("Fernando Habib", "Alejandro Anadon", "Santiago Gray", "Ignacio De La Vega"):
                for audience in ("profe", "atleta", "cliente"):
                    pdf = report_generator.generate_visual_report_pdf(state, athlete, audience)
                    cls.pdf[(athlete, audience)] = _pdf_text(pdf) if pdf else ""

    def test_rule1_imtp_quality_flag_in_dashboard_and_pdf(self):
        flag = _ascii(FLAG_TEXT)
        self.assertIn(flag, self.dashboard["Fernando Habib"])
        self.assertIn(flag, self.pdf[("Fernando Habib", "profe")])
        # Ningun valor de RFD de la toma marcada aparece en ninguna superficie.
        for text in (self.dashboard["Fernando Habib"], self.pdf[("Fernando Habib", "profe")]):
            self.assertNotIn("196 n/s", text)
            self.assertIsNone(re.search(r"rfd 50[^a-z]{0,6}196", text))

    def test_rule2_gross_and_net_relative_force_labelled(self):
        for text in (self.dashboard["Fernando Habib"], self.pdf[("Fernando Habib", "profe")]):
            self.assertIn("bruta", text)
            self.assertIn("neta", text)
            self.assertIn("43.03", text)
            self.assertIn("33.22", text)

    def test_rule3_zero_weight_has_no_relative_force(self):
        for text in (self.dashboard["Santiago Gray"], self.pdf[("Santiago Gray", "profe")]):
            self.assertIsNone(re.search(r"\binf\b", text))
            self.assertIsNone(re.search(r"\d+\.\d+\s*n/kg", text))

    def test_rule4_stale_label_in_dashboard_and_pdf(self):
        label = _ascii("Evaluación desactualizada (hace 25 semanas)")
        self.assertIn(label, self.dashboard["Alejandro Anadon"])
        self.assertIn(label, self.pdf[("Alejandro Anadon", "profe")])

    def test_rule5_threshold_source_label_in_dashboard_and_pdf(self):
        self.assertIn("1.2 (te instrumento 3%)", self.dashboard["Ignacio De La Vega"])
        self.assertNotIn("(fijo)", self.dashboard["Ignacio De La Vega"])
        self.assertIn("te instrumento 3%", self.pdf[("Ignacio De La Vega", "profe")])
        self.assertNotIn("(fijo)", self.pdf[("Ignacio De La Vega", "profe")])


if __name__ == "__main__":
    unittest.main()
