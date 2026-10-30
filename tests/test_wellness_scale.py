"""Una sola escala de estres/dolor en dashboard, PDF y Excel (0-10, TeamBuildr).

Defecto real (oct/2026): el PDF inferia la escala del maximo de los datos y
el mismo dolor salia como "/5" en un reporte y "/10" en otro.
"""

from __future__ import annotations

import io
import unittest

import pandas as pd
from openpyxl import load_workbook

from charts.load_charts import chart_weekly_wellness, chart_wellness
from modules import report_generator
from modules.data_loader import _wellness_score
from modules.wellness_scales import PAIN_LABEL, STRESS_LABEL, WELLNESS_RESPONSE_SCALE
from tests.test_weekly_acwr_context import _chart_theme


class WellnessScaleTest(unittest.TestCase):
    def test_scale_does_not_depend_on_data(self):
        low = pd.DataFrame({"Estres": [1, 2, 3], "Dolor": [0, 1, 2]})
        high = pd.DataFrame({"Estres": [6, 8, 9], "Dolor": [7, 8, 9]})
        for frame in (low, high):
            scales = report_generator._professional_wellness_scales(frame)
            self.assertEqual(scales["stress"], WELLNESS_RESPONSE_SCALE)
            self.assertEqual(scales["pain"], WELLNESS_RESPONSE_SCALE)
        self.assertEqual(WELLNESS_RESPONSE_SCALE, "/10")

    def test_score_uses_same_scale(self):
        self.assertEqual(_wellness_score(8, 0, 0), 28.0)
        self.assertEqual(_wellness_score(8, 10, 10), 8.0)

    def test_dashboard_charts_label_scale(self):
        daily = pd.DataFrame({"Date": pd.date_range("2026-10-01", periods=3), "Estres": [3, 4, 5], "Dolor": [1, 2, 3]})
        names = {trace.name for trace in chart_wellness(daily, "Ana", theme=_chart_theme()).data}
        self.assertIn(STRESS_LABEL, names)
        self.assertIn(PAIN_LABEL, names)
        weekly = pd.DataFrame(
            {"week_start": pd.date_range("2026-09-07", periods=3, freq="7D"), "Sueno_mean": [7, 7, 7], "Estres_mean": [3, 3, 3], "Dolor_mean": [2, 2, 2]}
        )
        names = {trace.name for trace in chart_weekly_wellness(weekly, "Ana", theme=_chart_theme()).data}
        self.assertTrue(any("(0-10)" in name for name in names))

    def test_excel_headers_carry_scale(self):
        wellness = pd.DataFrame(
            {"Athlete": ["Ana"], "Date": ["2026-10-01"], "Sueno_hs": [7.0], "Estres": [3], "Dolor": [2], "Wellness_Score": [22.0]}
        )
        workbook = load_workbook(io.BytesIO(report_generator.export_excel({"Wellness": wellness})))
        sheet = workbook[workbook.sheetnames[0]]
        headers = [cell.value for cell in sheet[1]]
        self.assertIn(STRESS_LABEL, headers)
        self.assertIn(PAIN_LABEL, headers)


if __name__ == "__main__":
    unittest.main()
